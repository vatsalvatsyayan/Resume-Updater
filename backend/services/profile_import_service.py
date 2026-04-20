from __future__ import annotations

import io
import json
import re

from fastapi import HTTPException, UploadFile, status
from pydantic import ValidationError

from schemas.profile_import import ImportedProfileData
from services import llm_client
from utils.date_import import normalize_imported_profile


class ProfileImportService:
    async def import_profile(
        self,
        resume_file: UploadFile | None = None,
        resume_text: str | None = None,
    ) -> tuple[ImportedProfileData, list[str], list[str]]:
        sources: list[str] = []
        text_source = (resume_text or "").strip()

        if resume_file is not None:
            text_source = await self._extract_text_from_upload(resume_file)
            sources.append("resume")
        elif text_source:
            sources.append("resume_text")

        if not text_source:
            return ImportedProfileData(), [], sources

        data = self._parse_resume_text_with_llm(text_source)
        warnings = []
        if not self._has_any_data(data):
            warnings.append(
                "The uploaded resume could be read, but very little structured data was extracted."
            )

        return data, warnings, sources

    async def _extract_text_from_upload(self, upload: UploadFile) -> str:
        file_bytes = await upload.read()
        if not file_bytes:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded resume file is empty.",
            )

        filename = (upload.filename or "resume").lower()
        extension = filename.rsplit(".", 1)[-1] if "." in filename else ""

        try:
            if extension == "pdf":
                return self._extract_pdf_text(file_bytes)
            if extension == "docx":
                return self._extract_docx_text(file_bytes)
            if extension in {"txt", "md"} or upload.content_type == "text/plain":
                return self._decode_text_bytes(file_bytes).strip()
        except RuntimeError as exc:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=str(exc),
            ) from exc
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Could not read the uploaded resume: {exc}",
            ) from exc

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Upload a .pdf, .docx, or .txt resume.",
        )

    def _extract_pdf_text(self, file_bytes: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pypdf is required to import PDF resumes.") from exc

        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()

    def _extract_docx_text(self, file_bytes: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is required to import DOCX resumes."
            ) from exc

        document = Document(io.BytesIO(file_bytes))
        parts: list[str] = []
        para_text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        if para_text.strip():
            parts.append(para_text.strip())
        # Tables often hold job titles + date ranges (paragraph-only extraction misses them).
        for table in document.tables:
            rows_out: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows_out.append(" | ".join(c for c in cells if c))
            if rows_out:
                parts.append("\n".join(rows_out))
        return "\n\n".join(parts).strip()

    def _decode_text_bytes(self, file_bytes: bytes) -> str:
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="ignore")

    def _parse_resume_text_with_llm(self, text: str) -> ImportedProfileData:
        raw = llm_client.generate(self._build_prompt(text))
        try:
            obj = json.loads(self._extract_json(raw))
            if isinstance(obj, dict) and isinstance(obj.get("data"), dict):
                obj = obj["data"]
            data = ImportedProfileData.model_validate(obj)
            data = normalize_imported_profile(data)
            return self._merge_contact_fallback(text, data)
        except (json.JSONDecodeError, ValidationError, TypeError) as exc:
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail=f"LLM did not return valid profile import data: {exc}",
            ) from exc

    def _merge_contact_fallback(
        self, resume_text: str, data: ImportedProfileData
    ) -> ImportedProfileData:
        """Fill phone from raw text when the model omits it but a number is visible."""
        pi = data.personalInfo
        if (pi.phone or "").strip():
            return data
        extracted = self._extract_phone_from_text(resume_text)
        if not extracted:
            return data
        return data.model_copy(
            update={"personalInfo": pi.model_copy(update={"phone": extracted})}
        )

    _PHONE_SEARCH_RE = re.compile(
        r"(?:\+?\d{1,3}[-.\s]?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|"
        r"\(?\d{3}\)?[-.\s]\d{3}[-.\s]\d{4}|"
        r"\b\d{3}[-.\s]\d{3}[-.\s]\d{4}\b)"
    )

    def _extract_phone_from_text(self, resume_text: str) -> str | None:
        # Prefer numbers on the same line as Phone/Mobile/Tel labels
        for line in resume_text.splitlines():
            if re.search(r"(?i)\b(phone|mobile|cell|tel\.?|telephone)\b", line):
                m = self._PHONE_SEARCH_RE.search(line)
                if m:
                    return self._normalize_phone_candidate(m.group(0))
        for m in self._PHONE_SEARCH_RE.finditer(resume_text):
            cand = self._normalize_phone_candidate(m.group(0))
            digits = re.sub(r"\D", "", cand)
            if len(digits) >= 10:
                return cand
        return None

    @staticmethod
    def _normalize_phone_candidate(raw: str) -> str:
        return re.sub(r"\s+", " ", raw.strip())

    def _build_prompt(self, text: str) -> str:
        return f"""Extract structured profile data from this resume for a user profile form.

Return only one valid JSON object. Do not include markdown or explanations.
Use empty strings, null, false, or [] for missing fields. Do not invent facts.

Contact and education (capture when present in the resume text):
- personalInfo.phone: digits and symbols exactly as shown (header, summary, or labeled "Phone").
- personalInfo.location: city and state/region or country as written for the candidate (not job office locations), e.g. "Los Angeles, CA".
- For each education entry: gpa must be filled when the resume lists GPA/CGPA (e.g. "3.8/4.0", "GPA: 3.95", "Major GPA 3.7").
- education.location: campus, city, or region of the school if stated (e.g. "Los Angeles" or "Bengaluru, India").

Dates (required whenever they appear in the resume text):
- For each education, work experience, volunteer, and leadership entry, set startDate and endDate from the resume. Parse ranges like "Jan 2020 – Mar 2023", "2020-2023", "01/2022 – Present", or years under job titles.
- Prefer ISO dates: always use YYYY-MM-DD. If only month+year is known, use the first day of that month (e.g. June 2021 -> 2021-06-01). If only a calendar year is known, use YYYY-01-01.
- If a role is ongoing, set isPresent to true, endDate to null, and startDate to the best-known start.
- For certifications, fill issueDate and expiryDate the same way when stated.
- Do not leave date fields null when the resume clearly states time ranges for that entry (even if PDF text is messy).

The JSON object must use exactly this shape:
{{
  "personalInfo": {{
    "name": "",
    "email": "",
    "phone": null,
    "location": null,
    "portfolioWebsite": null,
    "githubUrl": null,
    "linkedinUrl": null
  }},
  "education": [
    {{
      "universityName": "",
      "courseName": "",
      "courseType": "Bachelor's",
      "major": "",
      "gpa": null,
      "location": null,
      "startDate": null,
      "endDate": null,
      "isPresent": false
    }}
  ],
  "workExperience": [
    {{
      "companyName": "",
      "position": "",
      "location": null,
      "startDate": null,
      "endDate": null,
      "isPresent": false,
      "summary": null,
      "description": null
    }}
  ],
  "projects": [
    {{
      "projectName": "",
      "link": null,
      "techStack": [],
      "summary": null,
      "description": null
    }}
  ],
  "skills": {{
    "programmingLanguages": [],
    "frameworks": [],
    "databases": [],
    "toolsAndTechnologies": [],
    "cloud": [],
    "ai": [],
    "other": []
  }},
  "certifications": [
    {{
      "name": "",
      "issuingOrganization": "",
      "issueDate": null,
      "expiryDate": null,
      "hasNoExpiry": false,
      "credentialId": null,
      "credentialUrl": null
    }}
  ],
  "volunteer": [
    {{
      "organizationName": "",
      "role": "",
      "cause": null,
      "location": null,
      "startDate": null,
      "endDate": null,
      "isPresent": false,
      "description": null
    }}
  ],
  "leadership": [
    {{
      "title": "",
      "organization": "",
      "startDate": null,
      "endDate": null,
      "isPresent": false,
      "description": null
    }}
  ]
}}

For courseType, use one of: "Bachelor's", "Master's", "PhD", "Diploma", "Certificate", "Associate", "".
If a section has no entries, return an empty array for that section.

Resume text:
{text}"""

    def _extract_json(self, text: str) -> str:
        text = text.strip()
        match = re.search(r"```(?:json)?\s*([\s\S]*?)\s*```", text)
        return match.group(1).strip() if match else text

    def _has_any_data(self, data: ImportedProfileData) -> bool:
        personal_info = data.personalInfo
        return any(
            [
                personal_info.name,
                personal_info.email,
                personal_info.phone,
                personal_info.location,
                personal_info.portfolioWebsite,
                personal_info.githubUrl,
                personal_info.linkedinUrl,
                data.education,
                data.workExperience,
                data.projects,
                data.certifications,
                data.volunteer,
                data.leadership,
                *data.skills.model_dump().values(),
            ]
        )
